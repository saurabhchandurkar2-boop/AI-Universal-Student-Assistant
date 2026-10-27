"""
knowledge_loader.py - add uploaded PDF / Word / TXT / MD files to the knowledge base.

Uses your existing pipeline:
    documents_to_import -> DocumentIngester -> knowledge/documents -> KnowledgeManager

DocumentIngester already reads .txt, .md and .pdf. Word (.docx) is not supported
by it, so a .docx is converted to a .txt first (needs: pip install python-docx).
"""
import io
import os
import re
import sys
from pathlib import Path

from chunk_filter import remove_page_index, write_page_index

PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))
SOURCE_DIR = os.path.join(PROJECT_ROOT, "documents_to_import")
DOCS_DIR = os.path.join(PROJECT_ROOT, "knowledge", "documents")

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


def _safe_name(filename: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename).stem).strip("._")
    return stem or "document"


def _docx_to_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pdf_pages(data: bytes, progress=None):
    """Read a PDF page by page. Returns (page_texts, total_pages, pages_with_text).

    page_texts has one entry for every page (empty when a page has no text), so
    list position + 1 is the real page number.
    progress(done, total) is called after every page, so large PDFs can show
    a progress bar.
    """
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password protected.")

    total = len(reader.pages)
    texts, with_text = [], 0
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            with_text += 1
        texts.append(text)
        if progress:
            progress(number, total)
    return texts, total, with_text


def _pdf_to_text(data: bytes, progress=None):
    """Read a PDF page by page. Returns (text, total_pages, pages_with_text)."""
    texts, total, with_text = _pdf_pages(data, progress)
    return "\n\n".join(t for t in texts if t.strip()), total, with_text


def convert_upload(data: bytes, filename: str, progress=None, page_sink=None):
    """Turn an uploaded file into plain text.

    Returns (name, out_name, payload, pages, warning). Raises ValueError for
    unsupported or unreadable files. Used by add_to_knowledge and by the
    hosted per-visitor workspace.

    page_sink: optional list. For a PDF it receives the text of every page, so the
    caller can save a page index (see chunk_filter.write_page_index).
    """
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type '{ext}'. Use: {sorted(SUPPORTED)}")

    name = _safe_name(filename)
    pages, warning = None, ""

    if ext == ".docx":
        text = _docx_to_text(data).strip()
        if not text:
            raise ValueError(f"No readable text found in {filename}.")
        out_name, payload = name + ".txt", text.encode("utf-8")
    elif ext == ".pdf":
        page_texts, pages, with_text = _pdf_pages(data, progress)
        if page_sink is not None:
            page_sink.extend(page_texts)
        text = "\n\n".join(t for t in page_texts if t.strip()).strip()
        if not text:
            raise ValueError(
                f"No readable text found in {filename}. "
                "If it is a scanned PDF, it needs OCR first."
            )
        if with_text < pages * 0.5:
            warning = (
                f"Only {with_text} of {pages} pages contain text. "
                "Scanned pages need OCR first and were skipped."
            )
        out_name, payload = name + ".txt", text.encode("utf-8")
    else:
        out_name, payload = name + ext, data

    return name, out_name, payload, pages, warning


def add_to_knowledge(data: bytes, filename: str, progress=None) -> dict:
    """Save an uploaded file into documents_to_import and run the pipeline.

    Returns {"saved_as", "index_success", "pages", "warning"}.
    progress(done, total) is called per page for PDFs.
    Raises ValueError for unsupported or unreadable files.
    """
    os.makedirs(SOURCE_DIR, exist_ok=True)
    page_texts = []
    name, out_name, payload, pages, warning = convert_upload(
        data, filename, progress, page_sink=page_texts
    )

    with open(os.path.join(SOURCE_DIR, out_name), "wb") as f:
        f.write(payload)

    # Run your existing import + index pipeline.
    if PROJECT_ROOT not in sys.path:
        sys.path.insert(0, PROJECT_ROOT)
    from knowledge.ingest.document_ingester import DocumentIngester

    result = DocumentIngester().run_pipeline()

    # A scanned PDF has no text layer, so nothing lands in knowledge/documents.
    if not os.path.exists(os.path.join(DOCS_DIR, name + ".txt")):
        os.remove(os.path.join(SOURCE_DIR, out_name))
        raise ValueError(
            f"No readable text found in {filename}. "
            "If it is a scanned PDF, it needs OCR first."
        )

    if page_texts:
        write_page_index(DOCS_DIR, name, page_texts)

    return {
        "saved_as": name + ".txt",
        "index_success": result["index_success"],
        "pages": pages,
        "warning": warning,
    }


def list_documents() -> list:
    """Names of the documents currently in the knowledge base."""
    if not os.path.isdir(DOCS_DIR):
        return []
    found = []
    for root, _, files in os.walk(DOCS_DIR):
        for f in files:
            if f.lower().endswith(".txt"):
                found.append(os.path.relpath(os.path.join(root, f), DOCS_DIR))
    return sorted(found)


def read_document(rel_name: str) -> str:
    """Text of a document in the knowledge base."""
    target = os.path.abspath(os.path.join(DOCS_DIR, rel_name))
    if not target.startswith(os.path.abspath(DOCS_DIR) + os.sep) or not os.path.isfile(target):
        raise ValueError("Document not found.")
    with open(target, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def remove_document(rel_name: str) -> bool:
    """Remove a document from the knowledge base and rebuild the index.

    rel_name is a name returned by list_documents(). The matching source file
    in documents_to_import is removed too, otherwise the ingester would bring
    the document back on the next upload.
    """
    target = os.path.abspath(os.path.join(DOCS_DIR, rel_name))
    if not target.startswith(os.path.abspath(DOCS_DIR) + os.sep):
        raise ValueError("Invalid document name.")
    if not os.path.isfile(target):
        return False

    os.remove(target)

    stem = os.path.splitext(rel_name)[0]
    remove_page_index(os.path.dirname(target), os.path.splitext(os.path.basename(target))[0])
    for ext in (".txt", ".md", ".pdf"):
        source = os.path.abspath(os.path.join(SOURCE_DIR, stem + ext))
        if source.startswith(os.path.abspath(SOURCE_DIR) + os.sep) and os.path.isfile(source):
            os.remove(source)

    if PROJECT_ROOT not in sys.path:
        sys.path.insert(0, PROJECT_ROOT)
    from knowledge.knowledge_manager import KnowledgeManager

    manager = KnowledgeManager(documents_path=DOCS_DIR)
    manager.save_index(manager.build_index())
    return True


def ensure_initial_knowledge() -> int:
    """On first run, import the sample files in documents_to_import.

    Does nothing if knowledge/documents already has documents. Returns the
    number of documents in the knowledge base afterwards.
    """
    if list_documents():
        return len(list_documents())
    if not os.path.isdir(SOURCE_DIR) or not os.listdir(SOURCE_DIR):
        return 0
    if PROJECT_ROOT not in sys.path:
        sys.path.insert(0, PROJECT_ROOT)
    from knowledge.ingest.document_ingester import DocumentIngester

    DocumentIngester().run_pipeline()
    return len(list_documents())
